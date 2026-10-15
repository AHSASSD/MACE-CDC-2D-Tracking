"""按EvtID闭区间绘制完整事件与问题轨迹；复用原两级寻迹，不训练模型。"""
import argparse
from collections import Counter, defaultdict
import hashlib
import html
import json
from pathlib import Path
from datetime import datetime
import sys

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.collections import LineCollection
from matplotlib.lines import Line2D
from matplotlib.patches import Circle
import numpy as np
import pandas as pd
import torch

from common import ROOT, MAP_FILE, InteractionNetwork, model_dimensions
from STEP1_完整事件数据集 import event_graph
from STEP4_两级寻迹 import predict, run_two_pass, evaluate, truth_sets
from STEP4_困难事件分析 import best_track_match, classify_miss, graph_components

DATA = ROOT / 'whole_event_retraining'
ASSESS = DATA / 'stage2_no_time_assessment'
MODELS = {
    'last': (ASSESS / 'last_epoch_62.pth', ASSESS / 'probabilities_last.pt', '第二阶段62轮'),
    'best': (DATA / 'best_interaction_network_no_time_finetune_stage2_lr3e4.pth', ASSESS / 'probabilities_finetuned.pt', '第二阶段34轮'),
    'stage1': (DATA / 'best_interaction_network_no_time_finetune_whole_events_sigma05.pth', ASSESS / 'probabilities_baseline.pt', '第一阶段77轮'),
    'original': (ROOT / 'best_interaction_network_no_time_clean_physics_v2.pth', DATA / 'interrupted_no_time_assessment/probabilities_baseline.pt', '原始no-time'),
}
CAUSES = {'complete_but_contaminated': '候选完整但混入其他轨迹', 'edge_fragmentation_below_min_hits': '宽松连边仍碎成不足5hit',
          'strict_threshold_fragmentation': '严格阈值导致碎裂', 'merged_with_other_tracks': '连通块与其他真轨混合',
          'geometry_or_candidate_rejection': '几何或候选拒绝', 'partial_candidate_loss': '部分候选丢失'}
STATUS = {'matched': ('#149342', '匹配'), 'fake': ('#d62828', '假轨'), 'clone': ('#9b30b6', '克隆')}
SOURCE_COLUMNS = ['EvtID', 'TrkID', 'CellID', 'Good', 'WireType', 'CreatProc', 'd', 't', 'x', 'tHit']


def source_columns(columns):
    """Accept header casing/whitespace changes, but never guess missing physical fields."""
    names = {name.casefold(): name for name in SOURCE_COLUMNS + ['HitID']}
    renamed = [names.get(str(c).strip().lstrip('\ufeff').casefold(), str(c).strip()) for c in columns]
    if len(set(renamed)) != len(renamed):
        raise ValueError('表头包含重复列名')
    return renamed


def read_source_frames(path, sheet=None):
    if path.suffix.lower() == '.csv':
        header = pd.read_csv(path, nrows=0)
        names = source_columns(header.columns)
        missing = set(SOURCE_COLUMNS) - set(names)
        if missing:
            raise ValueError(f'{path.name} 缺少列：{sorted(missing)}')
        selected = [old for old, new in zip(header.columns, names) if new in SOURCE_COLUMNS + ['HitID']]
        offset = 2
        with pd.read_csv(path, usecols=selected, chunksize=50000) as reader:
            for frame in reader:
                frame.columns = source_columns(frame.columns)
                frame['source_row'] = np.arange(offset, offset + len(frame))
                offset += len(frame)
                yield frame
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            page = book[sheet] if sheet else book.worksheets[0]
            rows = page.iter_rows(values_only=True)
            names = source_columns(next(rows, ()))
            missing = set(SOURCE_COLUMNS) - set(names)
            if missing:
                raise ValueError(f'{path.name} / {page.title} 缺少列：{sorted(missing)}')
            indices = [i for i, name in enumerate(names) if name in SOURCE_COLUMNS + ['HitID']]
            fields = [names[i] for i in indices] + ['source_row']
            buffer = []
            for row_number, row in enumerate(rows, 2):
                if row_number % 100000 == 0:
                    print(f'{path.name}：已扫描 {row_number-1} 行…', flush=True)
                if not any(value is not None for value in row):
                    continue
                buffer.append([row[i] if i < len(row) else None for i in indices] + [row_number])
                if len(buffer) == 10000:
                    yield pd.DataFrame(buffer, columns=fields)
                    buffer = []
            if buffer:
                yield pd.DataFrame(buffer, columns=fields)
        finally:
            book.close()
    else:
        raise ValueError('数据文件仅支持 .xlsx 或 .csv')


def load_external_source(source, start, end, cell_map=MAP_FILE, sheet=None):
    if not 0 <= start <= end:
        raise ValueError('区间应满足0 <= 起点 <= 终点，两个端点均包含')
    source = Path(source).resolve()
    if source.is_dir():
        paths = sorted(p for p in source.iterdir() if p.is_file() and
                       p.suffix.lower() in {'.csv', '.xlsx'} and not p.name.startswith('~$'))
    elif source.is_file():
        paths = [source]
    else:
        raise FileNotFoundError(f'数据源不存在：{source}')
    frames, owners = [], {}
    for path in paths:
        print(f'读取 {path.name}，保留 EvtID {start}～{end}…', flush=True)
        try:
            for frame in read_source_frames(path, sheet):
                event_ids = pd.to_numeric(frame.EvtID, errors='raise')
                if event_ids.isna().any() or (event_ids % 1 != 0).any():
                    raise ValueError(f'{path.name} 中有空或非整数EvtID')
                selected = frame.loc[event_ids.between(start, end) &
                                     frame.Good.astype(str).str.strip().str.casefold().isin(['true', '1', '1.0']) &
                                     frame.WireType.astype(str).str.strip().eq('Axial') &
                                     frame.CreatProc.astype(str).str.strip().eq('ExtendedDecayWithSpin')].copy()
                if selected.empty:
                    continue
                for field in ['EvtID', 'TrkID', 'CellID', 'd', 't', 'tHit']:
                    selected[field] = pd.to_numeric(selected[field], errors='raise')
                if not np.isfinite(selected[['EvtID', 'TrkID', 'CellID', 'd', 't', 'tHit']].to_numpy(dtype=float)).all():
                    raise ValueError(f'{path.name} 中存在无效的ID、距离或时间')
                for field in ['EvtID', 'TrkID', 'CellID']:
                    if (selected[field] % 1 != 0).any():
                        raise ValueError(f'{path.name} 的{field}必须为整数')
                    selected[field] = selected[field].astype('int64')
                if 'HitID' not in selected:
                    # shortcut: 缺少HitID时使用原表行号标识，原始行号始终另存供追溯。
                    selected['HitID'] = selected.source_row
                selected['HitID'] = pd.to_numeric(selected.HitID, errors='raise')
                if not np.isfinite(selected.HitID).all() or (selected.HitID % 1 != 0).any():
                    raise ValueError(f'{path.name} 的HitID必须为有限整数')
                selected['HitID'] = selected.HitID.astype('int64')
                for evt in selected.EvtID.unique():
                    if int(evt) in owners and owners[int(evt)] != path:
                        raise ValueError(f'EvtID {evt} 同时出现在 {owners[int(evt)].name} 和 {path.name}；请单独选择文件，避免跨批次拼接')
                    owners[int(evt)] = path
                selected['source_file'] = str(path)
                frames.append(selected)
        except ValueError as exc:
            if source.is_dir() and '缺少列：' in str(exc):
                print(f'跳过非hit表：{exc}', flush=True)
                continue
            raise
    if not frames:
        raise ValueError('所选数据源和EvtID区间内没有满足 Good=True、Axial、ExtendedDecayWithSpin 的hit')
    geometry = pd.read_excel(cell_map).rename(columns={'x': 'wire_x', 'y': 'wire_y'})
    required = {'CellID', 'wire_x', 'wire_y', 'u', 'v'}
    if not required <= set(geometry.columns):
        raise ValueError(f'丝几何表缺少列：{sorted(required-set(geometry.columns))}')
    geometry = geometry.loc[geometry.u.eq(0) & geometry.v.eq(0), ['CellID', 'wire_x', 'wire_y']].copy()
    geometry['wire_r'] = np.hypot(geometry.wire_x, geometry.wire_y)
    combined = pd.concat(frames, ignore_index=True).merge(geometry, on='CellID', how='left', sort=False, validate='many_to_one')
    if not np.isfinite(combined[['wire_x', 'wire_y', 'wire_r']].to_numpy(dtype=float)).all():
        bad = combined.loc[combined.wire_x.isna() | combined.wire_y.isna(), 'CellID'].unique()
        raise ValueError(f'丝几何表缺少直丝CellID或坐标无效：{bad[:20].tolist()}')
    return {int(evt): frame.reset_index(drop=True) for evt, frame in combined.groupby('EvtID', sort=True)}


def select_ids(split, start, end, partition):
    if not 0 <= start <= end:
        raise ValueError('区间应满足0 <= 起点 <= 终点，两个端点均包含')
    groups = {'train': split['train_evt_ids'], 'validation': split['val_evt_ids'], 'test': split['test_evt_ids']}
    return {int(evt): name for name, ids in groups.items() if partition == 'all' or partition == name
            for evt in ids if start <= int(evt) <= end}


def load_graphs(partitions):
    result = {}
    source_indices = {evt // 1000 + 1 for evt in partitions}
    # shortcut: 当前数据按每1000个EvtID一个源文件分片；更换分片约定时改为读取清单索引。
    for name in sorted(set(partitions.values())):
        for path in sorted((DATA / f'dataset_{name}.pt.parts').glob('part_*.pt')):
            if int(path.name.split('_')[1]) not in source_indices:
                continue
            for graph in torch.load(path, map_location='cpu', weights_only=False):
                evt = int(graph.event_id.item())
                if evt in partitions:
                    if evt in result:
                        raise ValueError(f'重复EvtID：{evt}')
                    result[evt] = graph
    if set(result) != set(partitions):
        raise ValueError(f'缺少事件图：{sorted(set(partitions)-set(result))}')
    return result


def load_source(ids):
    frames = []
    columns = ['EvtID', 'TrkID', 'HitID', 'CellID', 'Good', 'WireType', 'CreatProc', 'd', 'x', 'tHit']
    for source_index in sorted({evt // 1000 + 1 for evt in ids}):
        paths = list((ROOT / 'whole data').glob(f'CDCSimHit_{source_index:02d}_*.csv'))
        if len(paths) != 1:
            raise ValueError(f'无法唯一定位第{source_index}个源CSV')
        for frame in pd.read_csv(paths[0], usecols=columns, chunksize=100000):
            selected = frame.loc[frame.EvtID.isin(ids) & frame.Good.eq(True) & frame.WireType.eq('Axial') & frame.CreatProc.eq('ExtendedDecayWithSpin')]
            if len(selected):
                frames.append(selected)
    if not frames:
        raise ValueError('区间内没有满足筛选条件的hit')
    return {int(evt): frame.reset_index(drop=True) for evt, frame in pd.concat(frames).groupby('EvtID', sort=True)}


def undirected_edges(graph, probability):
    edges = graph.edge_index.numpy().T
    pairs, inverse = np.unique(np.sort(edges, axis=1), axis=0, return_inverse=True)
    maxima = np.full(len(pairs), -np.inf)
    labels = np.zeros(len(pairs), dtype=bool)
    np.maximum.at(maxima, inverse, probability)
    np.logical_or.at(labels, inverse, graph.y.numpy().astype(bool))
    return pairs, maxima, labels


def circle_arc(fit, positions):
    a, b, radius = fit
    angles = np.sort(np.mod(np.arctan2(positions[:, 1]-b, positions[:, 0]-a), 2*np.pi))
    gaps = np.diff(np.r_[angles, angles[0]+2*np.pi])
    gap_index = int(np.argmax(gaps))
    padding = min(.08, 15/max(radius, 1))
    start = angles[(gap_index+1) % len(angles)]-padding
    span = min(2*np.pi, 2*np.pi-gaps[gap_index]+2*padding)
    theta = np.linspace(start, start+span, 240)
    return np.column_stack([a+radius*np.cos(theta), b+radius*np.sin(theta)])


def nearby_truth_segments(positions, max_gap_mm):
    pairs = np.stack([positions[:-1], positions[1:]], axis=1)
    return pairs[np.linalg.norm(positions[1:]-positions[:-1], axis=1) <= max_gap_mm]


def diagnose(graph, probability, strict, relaxed):
    tracks, diagnostics = run_two_pass(graph, None, None, 'no-time', strict, relaxed, probability)
    result, nodes, missed = evaluate_tracks(graph, probability, tracks, strict, relaxed)
    return tracks, result, nodes, missed, diagnostics


def evaluate_tracks(graph, probability, tracks, strict, relaxed):
    result = evaluate(graph, tracks)
    associations = truth_sets(graph)
    nodes = defaultdict(set)
    for index, keys in enumerate(associations):
        if len(keys) != 1:
            raise ValueError('当前绘图要求完整事件每hit有唯一真轨身份')
        for key in keys:
            nodes[key].add(index)
    seen = set()
    for track in tracks:
        key = tuple(track['truth'])
        valid = len(nodes[key]) >= 5 and track['purity'] >= .8 and track['completeness'] >= .5
        track['status'] = 'fake' if not valid else 'clone' if key in seen else 'matched'
        if valid:
            seen.add(key)
    assert Counter(t['status'] for t in tracks)['fake'] == result['fake_tracks']
    assert Counter(t['status'] for t in tracks)['clone'] == result['clone_tracks']
    sc, sm = graph_components(graph, probability, strict)
    rc, rm = graph_components(graph, probability, relaxed)
    missed = []
    for key in map(tuple, result['missed_tracks']):
        best = best_track_match(key, nodes[key], tracks, associations)
        cause, details = classify_miss(nodes[key], associations, sc, sm, rc, rm, best)
        missed.append({'truth': list(key), 'hits': len(nodes[key]), 'cause': CAUSES[cause], 'best_candidate': best, **details})
    return result, nodes, missed


def draw_event(evt, graph, frame, probability, tracks, result, nodes, missed, args, output):
    wire = np.column_stack([graph.wire_x.numpy(), graph.wire_y.numpy()])
    actual = np.array([[float(v) for v in str(value).strip('{}').split(',')][:2] for value in frame.x])
    if actual.shape != wire.shape or not np.isfinite(actual).all():
        raise ValueError(f'EvtID {evt}的模拟坐标x无效')
    keys = sorted(nodes)
    colors = {key: plt.get_cmap('turbo')((i+.5)/len(keys)) for i, key in enumerate(keys)}
    focus = next((key for key in keys if key[1] == args.track_id), None) if args.track_id is not None else (
        tuple(max(missed, key=lambda item: item['hits'])['truth']) if missed else max(keys, key=lambda key: len(nodes[key])))
    if focus is None:
        raise ValueError(f'EvtID {evt}没有TrkID {args.track_id}')
    members = nodes[focus]
    pairs, maxima, labels = undirected_edges(graph, probability)
    fig, axes = plt.subplots(2, 3, figsize=(19, 12))
    extent = max(np.abs(wire).max(), np.abs(actual).max()) + 30
    selected = np.array(sorted(members))
    local = np.vstack([wire[selected], actual[selected]])
    low, high = local.min(axis=0), local.max(axis=0)
    center = (low+high)/2
    half = max(float((high-low).max())/2+25, 40)
    for row in range(2):
        active = set(range(graph.num_nodes)) if row == 0 else members
        for ax in axes[row]:
            ax.scatter(*wire.T, s=7, c='#bcbcbc', alpha=.4, zorder=1)
            ax.set_aspect('equal')
            ax.set(xlabel='x (mm)', ylabel='y (mm)')
            ax.grid(alpha=.15)
            if row == 0:
                ax.set(xlim=(-extent, extent), ylim=(-extent, extent))
            else:
                ax.set(xlim=(center[0]-half, center[0]+half), ylim=(center[1]-half, center[1]+half))
        for key in keys if row == 0 else [focus]:
            indices = np.array(sorted(nodes[key]))
            ordered = indices[np.argsort(frame.tHit.to_numpy()[indices], kind='stable')]
            ax = axes[row, 0]
            ax.scatter(*wire[indices].T, color=colors[key], marker='x', s=18)
            ax.scatter(*actual[ordered].T, color=colors[key], s=9)
            ax.add_collection(LineCollection(nearby_truth_segments(actual[ordered], args.max_truth_gap_mm), colors=[colors[key]], linewidths=.8))
            ax.annotate(f'T{key[1]}', actual[ordered[-1]], fontsize=6, color=colors[key])
            if tuple(key) in set(map(tuple, result['missed_tracks'])):
                ax.scatter(*wire[indices].T, s=45, facecolors='none', edgecolors='#d62828', linewidths=1)
        ax = axes[row, 1]
        visible = np.ones(len(pairs), dtype=bool) if row == 0 else np.array([u in members or v in members for u, v in pairs])
        categories = [(labels & (maxima >= args.strict_threshold), '#169447', '-', '正确严格边'),
                      (labels & (maxima >= args.relaxed_threshold) & (maxima < args.strict_threshold), '#da9b00', '--', '正确宽松边'),
                      (~labels & (maxima >= args.relaxed_threshold), '#d62828', '-', '错误连边'),
                      (labels & (maxima < args.relaxed_threshold), '#777777', ':', '漏掉的真边')]
        for mask, color, style, label in categories:
            ax.add_collection(LineCollection(wire[pairs[mask & visible]], colors=color, linewidths=.65 if row else .35,
                                             linestyles=style, alpha=.7 if row else .3, label=label, zorder=2))
        ax.scatter(*wire[list(active)].T, c='#222222', s=9, zorder=3)
        ax.legend(loc='upper right', fontsize=7)
        ax = axes[row, 2]
        used = set()
        detail_handles = []
        for number, track in enumerate(tracks, 1):
            indices = track['indices']
            used.update(indices)
            if row == 1 and not members.intersection(indices):
                continue
            color, label = STATUS[track['status']]
            if track['status'] == 'matched':
                color = colors[tuple(track['truth'])]
            a, b, radius = track['fit']
            arc = circle_arc(track['fit'], wire[indices])
            style = {'matched': '-', 'fake': '--', 'clone': ':'}[track['status']]
            marker = {'matched': 'o', 'fake': 's', 'clone': 'D'}[track['status']]
            ax.plot(*arc.T, color=color, linewidth=1.3, linestyle=style, alpha=.8)
            ax.scatter(*wire[indices].T, color=color, marker=marker, s=18, zorder=4)
            ax.annotate(f'C{number}/T{track["truth"][1]}', wire[indices[0]], color=color, fontsize=6)
            if row:
                detail_handles.append(Line2D([], [], color=color, label=f'C{number}/T{track["truth"][1]} {label} | RMS={track["rms"]:.2f}mm\n纯度={track["purity"]:.2f} 完整度={track["completeness"]:.2f}'))
                for i in indices:
                    ax.add_patch(Circle(tuple(wire[i]), float(graph.drift_d[i]), fill=False, edgecolor=color, linewidth=.6, alpha=.7))
        unused = sorted(active-used)
        if unused:
            ax.scatter(*wire[unused].T, marker='x', c='#111111', s=30, zorder=5)
        status_handles = [Line2D([], [], color='#149342', label='匹配（与真值同色）'),
                          Line2D([], [], color=STATUS['fake'][0], linestyle='--', marker='s', label='假轨'),
                          Line2D([], [], color=STATUS['clone'][0], linestyle=':', marker='D', label='克隆')]
        ax.legend(handles=(detail_handles if row else status_handles)+[
            Line2D([], [], color='black', marker='x', linestyle='', label='未分配hit')], fontsize=7, loc='upper right')
    titles = ['真值：hit记录坐标 + 丝中心×；红圈表示漏检真轨', '原模型连边：双向概率最大值；红线跨真轨', '重建结果：拟合圆弧 + 使用的hit']
    for i, title in enumerate(titles):
        axes[0, i].set_title(title, fontsize=10)
        axes[1, i].set_title(f'局部放大：TrkID={focus[1]}，{len(members)} hit', fontsize=10)
    item = next((item for item in missed if tuple(item['truth']) == focus), None)
    reason = item['cause'] if item else '该真轨已匹配'
    recovery_label = {'none': '原流程', 'fragments': '时间约束片段回收', 'hough': '时间分峰Hough回收'}[getattr(args, 'recovery', 'none')]
    fig.suptitle(f'EvtID={evt} | {MODELS[args.model][2]} / {recovery_label} | {args.partitions[evt]} | 真轨{result["matched_tracks"]}/{result["truth_tracks"]}'
                 f' | 假轨{result["fake_tracks"]} | 克隆{result["clone_tracks"]} | 阈值{args.strict_threshold}/{args.relaxed_threshold}', fontsize=14)
    fig.text(.02, .025, f'局部观察T{focus[1]}：{reason}。hit记录坐标可能经过同Cell多步平均；仅连接间距≤{args.max_truth_gap_mm:g}mm的相邻点，不代表连续真实轨迹。\n'
             '漏检按整条轨迹匹配判断；未分配hit不一定等于漏检。拟合圆弧只显示hit覆盖角度，完整圆参数见JSON。局部连边包含相连的其他轨迹hit。', fontsize=10)
    fig.tight_layout(rect=(0, .07, 1, .96))
    fig.savefig(output, dpi=args.dpi)
    plt.close(fig)
    return focus


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--gui', action='store_true', help='打开交互式数据源与绘图参数窗口')
    parser.add_argument('--source', type=Path, help='外部.xlsx/.csv文件，或包含这些文件的目录；不指定则使用原缓存数据集')
    parser.add_argument('--cell-map', type=Path, default=MAP_FILE, help='外部数据构图使用的丝几何Excel表')
    parser.add_argument('--sheet', help='Excel工作表名称；默认第一个工作表')
    parser.add_argument('--event-range', nargs=2, type=int, metavar=('START', 'END'))
    parser.add_argument('--split', choices=['validation', 'train', 'test', 'all'], help='默认原数据用validation，外部数据用all')
    parser.add_argument('--model', choices=MODELS, default='last')
    parser.add_argument('--recovery', choices=['none', 'fragments', 'hough'], default='none', help='实验回收模式；默认仍是原流程，当前仅验证第62轮权重')
    parser.add_argument('--track-id', type=int, help='局部放大指定TrkID；不指定时优先放大最长漏检轨迹')
    parser.add_argument('--strict-threshold', type=float, default=.70)
    parser.add_argument('--relaxed-threshold', type=float, default=.575)
    parser.add_argument('--output-dir', type=Path)
    parser.add_argument('--dpi', type=int, default=150)
    parser.add_argument('--max-truth-gap-mm', type=float, default=50, help='仅绘图：超过该点距的真值记录不连线，保留所有hit')
    args = parser.parse_args()
    if args.gui or len(sys.argv) == 1:
        from STEP6_绘图界面 import launch
        launch()
        return
    if args.event_range is None:
        parser.error('请提供--event-range START END，或用--gui打开窗口')
    args.split = args.split or ('all' if args.source else 'validation')
    if args.source and args.split != 'all':
        parser.error('外部数据不沿用旧数据的训练/验证划分，请指定--split all')
    if args.recovery != 'none' and args.model != 'last':
        parser.error('当前回收参数仅验证第62轮权重，请使用--model last')
    if not 0 < args.relaxed_threshold <= args.strict_threshold < 1 or args.dpi < 50:
        parser.error('要求0 < 宽松阈值 <= 严格阈值 < 1，dpi至少50')
    if not np.isfinite(args.max_truth_gap_mm) or args.max_truth_gap_mm <= 0:
        parser.error('max-truth-gap-mm必须是有限正数')
    torch.set_num_threads(2)
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    start, end = args.event_range
    if args.source:
        source = load_external_source(args.source, start, end, args.cell_map, args.sheet)
        args.partitions = dict.fromkeys(source, 'external')
        graphs = {evt: event_graph(frame) for evt, frame in source.items()}
        print(f'外部数据：{len(graphs)}个事件，{sum(len(f) for f in source.values())}个筛选后hit；不加噪声、不重复加误差。', flush=True)
    else:
        split = json.loads((DATA / 'event_split.json').read_text(encoding='utf-8'))
        args.partitions = select_ids(split, start, end, args.split)
        graphs = load_graphs(args.partitions)
        source = load_source(set(graphs))
    if not args.partitions:
        parser.error('区间内没有所选数据组的事件，可扩大区间或指定--split all')
    suffix = '' if args.recovery == 'none' else f'_{args.recovery}'
    if args.source:
        suffix += '_external_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    output = args.output_dir or DATA / 'event_visualizations' / f'{args.model}_{args.split}_{start}_{end}{suffix}'
    if args.source and output.exists() and any(output.iterdir()):
        raise FileExistsError(f'外部数据的输出目录非空，请选择新目录以保留已有结果：{output}')
    output.mkdir(parents=True, exist_ok=True)
    weight, cache_path, model_label = MODELS[args.model]
    sha = hashlib.sha256(weight.read_bytes()).hexdigest()
    cached = torch.load(cache_path, map_location='cpu', weights_only=False) if not args.source and cache_path.exists() else None
    validation_parts = sorted((DATA / 'dataset_validation.pt.parts').glob('part_*.pt'))
    signature = {'weight_sha256': sha, 'validation_parts': [{'name': p.name, 'bytes': p.stat().st_size, 'mtime_ns': p.stat().st_mtime_ns} for p in validation_parts]}
    if cached and cached['signature'] != signature:
        print('概率缓存签名不匹配，将对选定事件重新推理。', flush=True)
        cached = None
    model = None
    rows = []
    for evt, graph in sorted(graphs.items()):
        frame = source[evt]
        expected = [(evt, int(t)) for t in frame.TrkID]
        if not np.array_equal(graph.cell_id.numpy(), frame.CellID.to_numpy()) or not np.allclose(graph.drift_d.numpy(), frame.d.to_numpy()) or truth_sets(graph) != [{key} for key in expected]:
            raise ValueError(f'EvtID {evt}源CSV与缓存图的hit身份/顺序不一致')
        if cached and evt in cached['values']:
            probability = cached['values'][evt]
        else:
            if model is None:
                model = InteractionNetwork(*model_dimensions('no-time'))
                model.load_state_dict(torch.load(weight, map_location='cpu', weights_only=True))
                model.eval()
            probability = predict(graph, model, torch.device('cpu'), 'no-time')
        if len(probability) != graph.num_edges or not np.isfinite(probability).all():
            raise ValueError('模型概率数量或数值无效')
        tracks, result, nodes, missed, diagnostics = diagnose(graph, probability, args.strict_threshold, args.relaxed_threshold)
        if args.recovery != 'none':
            from fragment_recovery import recover_fragments, hough_remaining
            tracks = [{k: v for k, v in t.items() if k not in ['truth', 'purity', 'completeness', 'status']} for t in tracks]
            tracks = recover_fragments(graph, tracks, probability, use_time=True)
            if args.recovery == 'hough':
                tracks = hough_remaining(graph, tracks, probability, use_time=True, seed=20261014+evt)
            result, nodes, missed = evaluate_tracks(graph, probability, tracks, args.strict_threshold, args.relaxed_threshold)
            diagnostics['recovery_profile'] = args.recovery
        png = f'event_{evt}.png'
        focus = draw_event(evt, graph, frame, probability, tracks, result, nodes, missed, args, output/png)
        circles = []
        def provenance(index):
            return {'source_file': str(frame.source_file.iloc[index]), 'source_row': int(frame.source_row.iloc[index])} if args.source else {}
        for number, track in enumerate(tracks, 1):
            circles.append({'circle_id': number, **track, 'source_hits': [
                {'node_index': i, 'EvtID': evt, 'TrkID': int(frame.TrkID.iloc[i]), 'HitID': int(frame.HitID.iloc[i]),
                 'CellID': int(frame.CellID.iloc[i]), 'd_mm': float(graph.drift_d[i]), **provenance(i)} for i in track['indices']]})
        record = {'EvtID': evt, 'partition': args.partitions[evt], 'model': model_label, 'weight_sha256': sha,
                  'thresholds': [args.strict_threshold, args.relaxed_threshold], 'focus_truth': list(focus),
                  'result': result, 'miss_diagnosis': missed, 'circles': circles, 'diagnostics': diagnostics}
        record['plot_max_truth_gap_mm'] = args.max_truth_gap_mm
        record['recovery_profile'] = args.recovery
        record['source'] = str(args.source.resolve()) if args.source else 'original_dataset'
        if args.source:
            record['cell_map'] = str(args.cell_map.resolve())
            record['sheet'] = args.sheet or '(first worksheet)'
        used = {i for track in tracks for i in track['indices']}
        record['hits'] = [{'node_index': i, 'TrkID': int(frame.TrkID.iloc[i]), 'HitID': int(frame.HitID.iloc[i]),
                           'CellID': int(frame.CellID.iloc[i]), 'wire_x_mm': float(graph.wire_x[i]),
                           'wire_y_mm': float(graph.wire_y[i]), 'd_mm': float(graph.drift_d[i]),
                           'unassigned': i not in used, **provenance(i)} for i in range(graph.num_nodes)]
        (output/f'event_{evt}.json').write_text(json.dumps(record, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
        rows.append({'EvtID': evt, 'partition': args.partitions[evt], 'png': png, 'focus_TrkID': focus[1], **result})
        print(f'EvtID={evt} [{args.partitions[evt]}] 真轨={result["matched_tracks"]}/{result["truth_tracks"]}，假轨={result["fake_tracks"]}，克隆={result["clone_tracks"]} -> {png}', flush=True)
    (output/'summary.json').write_text(json.dumps({'event_range_inclusive': [start, end], 'split': args.split, 'model': model_label, 'recovery': args.recovery, 'source': str(args.source.resolve()) if args.source else 'original_dataset', 'events': rows}, ensure_ascii=False, indent=2), encoding='utf-8')
    links = ''.join(f'<a href="#e{r["EvtID"]}">{r["EvtID"]}</a> ' for r in rows)
    cards = ''.join(f'<article id="e{r["EvtID"]}"><h2>EvtID {r["EvtID"]} · {html.escape(r["partition"])} · 找回{r["matched_tracks"]}/{r["truth_tracks"]} · 假轨{r["fake_tracks"]} · 克隆{r["clone_tracks"]}</h2>'
                    f'<a href="{r["png"]}" target="_blank"><img loading="lazy" src="{r["png"]}" alt="事件{r["EvtID"]}诊断图"></a><p><a href="event_{r["EvtID"]}.json">圆参数、hit归属与漏检分析</a> · <a href="#top">返回事件目录</a></p></article>' for r in rows)
    page = f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>事件切片诊断</title><style>body{{font:16px sans-serif;max-width:1800px;margin:24px auto;padding:0 20px;background:#f4f5f7}}img{{width:100%}}article{{background:white;padding:16px;margin:24px 0}}a{{color:#175bb4}}nav a{{display:inline-block;padding:6px}}</style><h1 id="top">{html.escape(model_label)} · EvtID闭区间{start}～{end}</h1><p>本页{len(rows)}个{args.split}事件。点击图像查看原尺寸；每事件下排自动放大一条问题真轨。JSON中fit=[圆心x,圆心y,半径]，单位mm。</p><nav>{links}</nav>{cards}</html>'
    (output/'index.html').write_text(page, encoding='utf-8')
    print(f'完成，浏览目录：{output / "index.html"}', flush=True)


if __name__ == '__main__':
    main()
